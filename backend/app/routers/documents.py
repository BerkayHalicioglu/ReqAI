import os
import uuid
import pypdf
import docx

from fastapi import APIRouter, Depends, UploadFile, File, HTTPException, BackgroundTasks
from fastapi.responses import FileResponse
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from app.database import get_db
from app.core.config import settings
from app.models.entities import Document, DocumentStatus, Requirement, Task, TestScenario, TestStatus
from app.schemas.schemas import (
    DocumentSummaryOut, DocumentDetailOut, RequirementUpdate, RequirementOut, 
    DocumentUpdate, PaginatedDocumentsOut, AnalyticsOut, TestScenarioOut, TestScenarioUpdate
)
from app.services.analysis_service import run_analysis, run_analysis_background
from app.services.ai_service import get_ai_service

router = APIRouter(prefix="/api/documents", tags=["documents"])


def extract_text_from_file(file: UploadFile) -> str:
    filename = file.filename.lower()
    if filename.endswith(".txt"):
        raw_bytes = file.file.read()
        try:
            return raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return raw_bytes.decode("latin-1")
    elif filename.endswith(".pdf"):
        try:
            reader = pypdf.PdfReader(file.file)
            text = ""
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            return text
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read PDF: {str(e)}")
    elif filename.endswith(".docx"):
        try:
            doc = docx.Document(file.file)
            text = ""
            for para in doc.paragraphs:
                text += para.text + "\n"
            return text
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read DOCX: {str(e)}")
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format.")


@router.post("/upload", response_model=DocumentSummaryOut, status_code=201)
def upload_document(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Step 1 & 2: upload a TXT, PDF, or DOCX requirements document and parse its content."""
    filename_lower = file.filename.lower()
    if not (filename_lower.endswith(".txt") or filename_lower.endswith(".pdf") or filename_lower.endswith(".docx")):
        raise HTTPException(status_code=400, detail="Only .txt, .pdf, and .docx files are supported.")

    text = extract_text_from_file(file)
    if not text.strip():
        raise HTTPException(status_code=400, detail="Document content is empty or could not be extracted.")

    # Persist a copy on disk too (handy for audit/debugging), keyed by a UUID.
    stored_name = f"{uuid.uuid4()}_{file.filename}"
    stored_path = os.path.join(settings.UPLOAD_DIR, stored_name)
    with open(stored_path, "w", encoding="utf-8") as f:
        f.write(text)

    document = Document(
        filename=file.filename,
        content=text,
        status=DocumentStatus.UPLOADED,
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


@router.post("/upload-evidence", status_code=201)
def upload_evidence(file: UploadFile = File(...)):
    """Upload an image or file evidence attachment to /uploads directory."""
    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
    stored_filename = f"evidence_{uuid.uuid4().hex[:8]}_{file.filename}"
    stored_path = os.path.join(settings.UPLOAD_DIR, stored_filename)

    with open(stored_path, "wb") as f:
        f.write(file.file.read())

    return {
        "filename": stored_filename,
        "original_filename": file.filename,
        "url": f"http://localhost:8000/api/documents/uploads/{stored_filename}"
    }


@router.get("/uploads/{filename}")
def get_uploaded_file(filename: str):
    """Retrieve and serve an uploaded evidence file or image."""
    file_path = os.path.join(settings.UPLOAD_DIR, filename)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="Evidence file not found.")
    return FileResponse(file_path)



@router.post("/{document_id}/analyze", response_model=DocumentSummaryOut)
def analyze_document(
    document_id: str,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db)
):
    """Step 3 & 4: trigger AI analysis in the background and return immediately."""
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")

    if document.status != DocumentStatus.PROCESSING:
        document.status = DocumentStatus.PROCESSING
        db.commit()
        db.refresh(document)

    background_tasks.add_task(run_analysis_background, document_id)
    return document



@router.get("", response_model=PaginatedDocumentsOut)
def list_documents(
    skip: int = 0,
    limit: int = 5,
    db: Session = Depends(get_db)
):
    """Step 5: list previously uploaded/analyzed documents with backend pagination."""
    total = db.query(Document).count()
    total_requirements = db.query(Requirement).count()
    total_critical_requirements = db.query(Requirement).filter(
        Requirement.priority.in_(["CRITICAL", "HIGH"])
    ).count()
    total_tasks = db.query(Task).count()

    items = (
        db.query(Document)
        .order_by(Document.uploaded_at.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return {
        "items": items,
        "total": total,
        "total_requirements": total_requirements,
        "total_critical_requirements": total_critical_requirements,
        "total_tasks": total_tasks
    }


@router.get("/analytics", response_model=AnalyticsOut)
def get_analytics(db: Session = Depends(get_db)):
    """Calculate and return global system requirement analytics and stats."""
    total_docs = db.query(Document).count()
    total_reqs = db.query(Requirement).count()
    total_tsks = db.query(Task).count()
    total_tests = db.query(TestScenario).count()
    
    # Requirements by Priority
    req_pri_counts = (
        db.query(Requirement.priority, func.count(Requirement.id))
        .group_by(Requirement.priority)
        .all()
    )
    req_pri_dict = {p.value if hasattr(p, "value") else str(p): count for p, count in req_pri_counts}
    for p in ["LOW", "MEDIUM", "HIGH", "CRITICAL"]:
        if p not in req_pri_dict:
            req_pri_dict[p] = 0
            
    # Tasks by Complexity
    tsk_cpl_counts = (
        db.query(Task.complexity, func.count(Task.id))
        .group_by(Task.complexity)
        .all()
    )
    tsk_cpl_dict = {c.value if hasattr(c, "value") else str(c): count for c, count in tsk_cpl_counts}
    for c in ["SIMPLE", "MODERATE", "COMPLEX"]:
        if c not in tsk_cpl_dict:
            tsk_cpl_dict[c] = 0
            
    return {
        "total_documents": total_docs,
        "total_requirements": total_reqs,
        "total_tasks": total_tsks,
        "total_test_scenarios": total_tests,
        "requirements_by_priority": req_pri_dict,
        "tasks_by_complexity": tsk_cpl_dict
    }


@router.get("/{document_id}", response_model=DocumentDetailOut)
def get_document(document_id: str, db: Session = Depends(get_db)):
    """Step 5: view a single document's full requirement/task/test breakdown."""
    document = (
        db.query(Document)
        .options(
            joinedload(Document.requirements)
            .joinedload(Requirement.tasks)
            .joinedload(Task.test_scenarios)
        )
        .filter(Document.id == document_id)
        .first()
    )
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")
    return document


@router.patch("/{document_id}/requirements/{requirement_id}", response_model=RequirementOut)
def update_requirement(
    document_id: str,
    requirement_id: str,
    payload: RequirementUpdate,
    db: Session = Depends(get_db)
):
    """Update a specific business requirement's title, description, or priority."""
    requirement = (
        db.query(Requirement)
        .filter(Requirement.document_id == document_id, Requirement.id == requirement_id)
        .first()
    )
    if not requirement:
        raise HTTPException(status_code=404, detail="Requirement not found.")

    requirement.title = payload.title
    requirement.description = payload.description
    requirement.priority = payload.priority

    db.commit()
    db.refresh(requirement)
    return requirement


@router.patch("/{document_id}", response_model=DocumentSummaryOut)
def update_document(
    document_id: str,
    payload: DocumentUpdate,
    db: Session = Depends(get_db)
):
    """Update document properties, such as renaming its filename."""
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")

    document.filename = payload.filename
    db.commit()
    db.refresh(document)
    return document


@router.delete("/{document_id}", status_code=204)
def delete_document(document_id: str, db: Session = Depends(get_db)):
    """Delete a document and all its cascaded requirements/tasks/test scenarios."""
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")

    db.delete(document)
    db.commit()
    return None


@router.patch("/{document_id}/test-scenarios/{scenario_id}", response_model=TestScenarioOut)
def update_test_scenario(
    document_id: str,
    scenario_id: str,
    payload: TestScenarioUpdate,
    db: Session = Depends(get_db)
):
    """Update a test scenario's status, comment, or evidence file."""
    scenario = db.query(TestScenario).filter(TestScenario.id == scenario_id).first()
    if not scenario:
        raise HTTPException(status_code=404, detail="Test scenario not found.")

    if payload.status is not None:
        scenario.status = payload.status
    if payload.comment is not None:
        scenario.comment = payload.comment
    if payload.evidence_file is not None:
        scenario.evidence_file = payload.evidence_file

    db.commit()
    db.refresh(scenario)
    return scenario

